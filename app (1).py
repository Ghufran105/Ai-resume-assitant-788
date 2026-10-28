"""
ATS Resume Checker
Upload a resume (PDF / DOCX / TXT) -> get an ATS score + concrete improvements.
UI: Streamlit | AI: Groq
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# Groq retires models from time to time. If one stops working, change it here
# or in the sidebar. Check https://console.groq.com/docs/models for current ones.
MODEL_OPTIONS = [
    "openai/gpt-oss-120b",  # best quality (default)
    "openai/gpt-oss-20b",   # faster / cheaper
    "qwen/qwen3.6-27b",     # alternative
]
MAX_RESUME_CHARS = 12000   # keep the prompt small enough for free-tier limits
MAX_JD_CHARS = 5000

# Weights for the final ATS score (must add up to 1.0)
WEIGHTS = {
    "keywords": 0.30,
    "formatting": 0.20,
    "content_quality": 0.25,
    "structure": 0.15,
    "readability": 0.10,
}

SYSTEM_PROMPT = """You are a strict, experienced ATS (Applicant Tracking System) \
analyst and professional resume reviewer. You will be given the plain text of a \
resume and optionally a job description.

Score the resume honestly. Do NOT inflate scores. A typical average resume \
should land between 50 and 70. Only give 85+ to truly excellent resumes.

Return ONLY a valid JSON object (no markdown, no commentary) with exactly this shape:
{
  "scores": {
    "keywords": <0-100, relevance and coverage of industry/job keywords>,
    "formatting": <0-100, ATS-friendly layout, consistent dates, no odd symbols>,
    "content_quality": <0-100, quantified achievements, strong action verbs>,
    "structure": <0-100, standard sections present and in a logical order>,
    "readability": <0-100, clarity, concise bullets, no fluff, no typos>
  },
  "target_role_guess": "<best guess of the role this resume targets>",
  "summary": "<2-3 sentence honest overall assessment>",
  "strengths": ["<short point>", "..."],
  "improvements": [
    {"priority": "High|Medium|Low", "issue": "<what is wrong>", "fix": "<exactly how to fix it>"}
  ],
  "missing_keywords": ["<keyword that should likely appear>", "..."],
  "bullet_rewrites": [
    {"original": "<weak bullet copied from resume>", "improved": "<stronger rewritten bullet>"}
  ]
}

Rules:
- Give 3-6 strengths, 5-10 improvements (sorted High to Low), up to 12 missing \
keywords, and up to 4 bullet rewrites.
- Only rewrite bullets that actually exist in the resume. Never invent \
employers, degrees, numbers or skills the candidate does not have; when \
suggesting a metric, use a placeholder like [X%] so the candidate fills in the real value.
- If a job description is provided, base 'keywords' and 'missing_keywords' on it. \
Otherwise base them on the target role you infer.
- Scores must be integers."""


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Return plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # resumes often put content inside tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# Quick rule-based checks (no AI needed)
# ----------------------------------------------------------------------------
SECTION_PATTERNS = {
    "Contact info": r"@|\+?\d[\d\s\-()]{8,}",
    "Summary / Objective": r"\b(summary|objective|profile|about me)\b",
    "Experience": r"\b(experience|employment|work history|internship)s?\b",
    "Education": r"\b(education|academic|university|college|degree|bachelor|master)\b",
    "Skills": r"\b(skills|technologies|technical skills|tools)\b",
    "Projects": r"\b(projects?)\b",
    "Certifications": r"\b(certifications?|certificates?|courses?)\b",
}


def basic_checks(text: str) -> dict:
    lower = text.lower()
    words = re.findall(r"\b\w+\b", text)
    sections = {k: bool(re.search(p, lower)) for k, p in SECTION_PATTERNS.items()}
    return {
        "word_count": len(words),
        "has_email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
        "has_phone": bool(re.search(r"\+?\d[\d\s\-()]{8,}\d", text)),
        "has_linkedin": "linkedin.com" in lower,
        "has_numbers": len(re.findall(r"\d+\s?%|\$\s?\d|\b\d{2,}\b", text)),
        "sections": sections,
    }


# ----------------------------------------------------------------------------
# Groq call
# ----------------------------------------------------------------------------
def _parse_json(raw: str) -> dict:
    """Parse JSON even if the model wrapped it in code fences or extra text."""
    raw = raw.strip()
    raw = re.sub(r"^```(?:json)?|```$", "", raw, flags=re.MULTILINE).strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def clean_result(data: dict) -> dict:
    """Make sure the model output has every field we need, with safe types."""
    scores_in = data.get("scores") or {}
    scores = {k: _clamp(scores_in.get(k)) for k in WEIGHTS}
    overall = round(sum(scores[k] * w for k, w in WEIGHTS.items()))

    def str_list(x):
        return [str(i) for i in x if str(i).strip()] if isinstance(x, list) else []

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).title(),
                    "issue": str(item.get("issue", "")),
                    "fix": str(item.get("fix", "")),
                }
            )
        elif isinstance(item, str):
            improvements.append({"priority": "Medium", "issue": item, "fix": ""})

    rewrites = []
    for item in data.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append(
                {"original": str(item["original"]), "improved": str(item["improved"])}
            )

    return {
        "overall": overall,
        "scores": scores,
        "role": str(data.get("target_role_guess", "")),
        "summary": str(data.get("summary", "")),
        "strengths": str_list(data.get("strengths")),
        "improvements": improvements,
        "missing_keywords": str_list(data.get("missing_keywords")),
        "bullet_rewrites": rewrites,
    }


def analyze_resume(client: Groq, model: str, resume_text: str, job_desc: str) -> dict:
    user_msg = f"RESUME TEXT:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\"\n"
    if job_desc.strip():
        user_msg += f"\nJOB DESCRIPTION:\n\"\"\"\n{job_desc[:MAX_JD_CHARS]}\n\"\"\"\n"
    else:
        user_msg += "\nNo job description provided. Infer the target role from the resume.\n"

    last_error = None
    for _ in range(2):  # one retry if the model returns broken JSON
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": user_msg},
            ],
            temperature=0.2,
            max_completion_tokens=6000,
            response_format={"type": "json_object"},
        )
        content = response.choices[0].message.content or ""
        try:
            return clean_result(_parse_json(content))
        except (json.JSONDecodeError, AttributeError) as err:
            last_error = err
    raise ValueError(f"The AI returned an unreadable answer. Please try again. ({last_error})")


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def score_label(score: int) -> tuple[str, str]:
    if score >= 80:
        return "Excellent", "🟢"
    if score >= 65:
        return "Good", "🟡"
    if score >= 50:
        return "Needs work", "🟠"
    return "Weak", "🔴"


def build_report(result: dict, checks: dict) -> str:
    lines = [
        "ATS RESUME REPORT",
        "=" * 40,
        f"Overall ATS score: {result['overall']}/100",
        f"Target role (guess): {result['role']}",
        "",
        "Score breakdown:",
    ]
    lines += [f"  - {k.replace('_', ' ').title()}: {v}/100" for k, v in result["scores"].items()]
    lines += ["", "Summary:", result["summary"], "", "Strengths:"]
    lines += [f"  + {s}" for s in result["strengths"]]
    lines += ["", "Improvements:"]
    for i, imp in enumerate(result["improvements"], 1):
        lines.append(f"  {i}. [{imp['priority']}] {imp['issue']}")
        if imp["fix"]:
            lines.append(f"     Fix: {imp['fix']}")
    lines += ["", "Missing keywords: " + ", ".join(result["missing_keywords"])]
    if result["bullet_rewrites"]:
        lines += ["", "Bullet rewrites:"]
        for r in result["bullet_rewrites"]:
            lines += [f"  Before: {r['original']}", f"  After:  {r['improved']}", ""]
    return "\n".join(lines)


def get_api_key() -> str:
    """Key priority: Streamlit secrets -> environment variable -> sidebar input."""
    try:
        if "GROQ_API_KEY" in st.secrets:
            return st.secrets["GROQ_API_KEY"]
    except Exception:
        pass  # no secrets file, that's fine
    return os.environ.get("GROQ_API_KEY", "")


# ----------------------------------------------------------------------------
# App
# ----------------------------------------------------------------------------
def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume and get an ATS score with clear fixes. Powered by Groq.")

    with st.sidebar:
        st.header("⚙️ Settings")
        api_key = get_api_key()
        if api_key:
            st.success("Groq API key loaded.")
        else:
            api_key = st.text_input("Groq API key", type="password", help="Get one free at console.groq.com")
        model = st.selectbox("Model", MODEL_OPTIONS, index=0)
        st.markdown("---")
        st.caption("Your resume is sent to Groq for analysis and is not stored by this app.")

    col_a, col_b = st.columns(2)
    with col_a:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col_b:
        job_desc = st.text_area(
            "Job description (optional, but gives a much better keyword score)",
            height=150,
            placeholder="Paste the job posting here...",
        )

    if st.button("🔍 Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please enter your Groq API key in the sidebar.")
            st.stop()

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
        except Exception as err:
            st.error(f"Could not read the file: {err}")
            st.stop()

        if len(text) < 100:
            st.error(
                "Almost no text could be extracted. If your resume is a scanned image or "
                "built from images, ATS systems cannot read it either. "
                "Export it again as a text-based PDF or DOCX."
            )
            st.stop()

        try:
            with st.spinner("Analyzing with Groq..."):
                client = Groq(api_key=api_key)
                result = analyze_resume(client, model, text, job_desc)
        except Exception as err:
            name = type(err).__name__
            if name == "AuthenticationError":
                st.error("Invalid Groq API key. Please check it and try again.")
            elif name == "RateLimitError":
                st.error("Groq rate limit reached. Wait a minute or switch to the 20b model.")
            elif name in ("NotFoundError", "BadRequestError"):
                st.error(
                    f"The model '{model}' was rejected by Groq. It may have been retired. "
                    f"Pick another model in the sidebar. Details: {err}"
                )
            else:
                st.error(f"Something went wrong: {err}")
            st.stop()

        st.session_state["result"] = result
        st.session_state["checks"] = basic_checks(text)

    # ---------------- Results ----------------
    result = st.session_state.get("result")
    checks = st.session_state.get("checks")
    if not result:
        st.info("Upload a resume and click **Analyze resume** to begin.")
        return

    label, icon = score_label(result["overall"])
    st.divider()
    top1, top2 = st.columns([1, 2])
    with top1:
        st.metric("ATS Score", f"{result['overall']} / 100")
        st.progress(result["overall"] / 100)
        st.subheader(f"{icon} {label}")
        if result["role"]:
            st.caption(f"Detected target role: **{result['role']}**")
    with top2:
        st.markdown("**Score breakdown**")
        for key, weight in WEIGHTS.items():
            val = result["scores"][key]
            st.write(f"{key.replace('_', ' ').title()} ({int(weight * 100)}% weight): **{val}**")
            st.progress(val / 100)

    if result["summary"]:
        st.info(result["summary"])

    tab1, tab2, tab3, tab4, tab5 = st.tabs(
        ["🛠️ Improvements", "💪 Strengths", "🔑 Keywords", "✍️ Bullet rewrites", "✅ Quick checks"]
    )

    with tab1:
        colors = {"High": "🔴", "Medium": "🟠", "Low": "🟡"}
        if not result["improvements"]:
            st.write("No improvements returned.")
        for imp in result["improvements"]:
            st.markdown(f"{colors.get(imp['priority'], '🟠')} **{imp['priority']}: {imp['issue']}**")
            if imp["fix"]:
                st.markdown(f"&nbsp;&nbsp;&nbsp;↳ *Fix:* {imp['fix']}")

    with tab2:
        for s in result["strengths"] or ["No strengths returned."]:
            st.markdown(f"- ✅ {s}")

    with tab3:
        if result["missing_keywords"]:
            st.write("Consider adding these (only if they are true for you):")
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("No missing keywords found.")

    with tab4:
        if not result["bullet_rewrites"]:
            st.write("No rewrites returned.")
        for r in result["bullet_rewrites"]:
            st.markdown(f"**Before:** {r['original']}")
            st.markdown(f"**After:** {r['improved']}")
            st.markdown("---")
        if result["bullet_rewrites"]:
            st.caption("Replace [X] placeholders with your real numbers. Never add facts that are not true.")

    with tab5:
        st.write(f"Word count: **{checks['word_count']}** (ideal: 400-800 for most resumes)")
        st.write(f"Email found: {'✅' if checks['has_email'] else '❌'}")
        st.write(f"Phone found: {'✅' if checks['has_phone'] else '❌'}")
        st.write(f"LinkedIn link: {'✅' if checks['has_linkedin'] else '❌'}")
        st.write(f"Numbers/metrics used: **{checks['has_numbers']}**")
        st.markdown("**Sections detected**")
        for sec, found in checks["sections"].items():
            st.write(f"{'✅' if found else '❌'} {sec}")

    st.download_button(
        "⬇️ Download report (.txt)",
        data=build_report(result, checks),
        file_name="ats_report.txt",
        mime="text/plain",
    )


if __name__ == "__main__":
    main()
